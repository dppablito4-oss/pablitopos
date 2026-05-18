import tkinter as tk
from tkinter import filedialog
import ttkbootstrap as tb
from .ui_sales_helpers import SalesUIHelpers
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

class ProductUIManager:
    def __init__(self, app):
        self.app = app

    # -------------------------------------------------------------------------
    # CATÁLOGO DE PRODUCTOS (PESTAÑA VENTAS)
    # -------------------------------------------------------------------------

    def load_products_for_sales(self):
        """Carga todos los productos disponibles en el catálogo de ventas."""
        try:
            # Obtener todos los productos activos
            products = self.app.product_service.search_products("")  # Buscar sin filtro = todos
            self.app.sales_products_cache = products
            
            # Usar Helper para llenar la tabla con indicadores de stock
            SalesUIHelpers.add_stock_indicator_to_product_list(
                self.app.tree_sales_products,
                products,
                self.app.product_service,
                reserved_items=getattr(self.app, "sale_items", [])
            )
            
        except Exception as e:
            print(f"Error cargando productos para ventas: {e}")
            self.app.sales_products_cache = []

    def filter_sales_products(self, event=None):
        """Filtra la tabla de productos según el texto de búsqueda."""
        try:
            search_text = self.app.entry_product_catalog_search.get().strip().upper()
            
            filtered_products = []
            if not search_text:
                filtered_products = self.app.sales_products_cache
            else:
                # Filtrar productos que coincidan con el texto
                for product in self.app.sales_products_cache:
                    # product: (id, code, name, unit, price, stock, price_includes_igv)
                    code = product[1]
                    name = product[2]
                    
                    if search_text in name.upper() or search_text in (code or "").upper():
                        filtered_products.append(product)
            
            # Usar Helper para llenar la tabla
            SalesUIHelpers.add_stock_indicator_to_product_list(
                self.app.tree_sales_products,
                filtered_products,
                self.app.product_service,
                reserved_items=getattr(self.app, "sale_items", [])
            )
            
        except Exception as e:
            print(f"Error filtrando productos: {e}")

    def on_sales_product_select(self, event=None):
        """Cuando se selecciona un producto del catálogo, prepararlo para agregar."""
        try:
            selection = self.app.tree_sales_products.focus()
            if not selection:
                return
            
            # El iid es el product_id
            product_id = int(selection)
            
            # Buscar el producto completo en el cache
            product = next((p for p in self.app.sales_products_cache if p[0] == product_id), None)
            
            if product:
                # product: (id, code, name, unit, price, stock, price_includes_igv)
                product_id = product[0]
                name = product[2]
                unit = product[3]
                price = product[4]
                
                # Guardar referencia del producto seleccionado
                self.app.selected_product = {
                    "id": product_id,
                    "name": name,
                    "unit": unit,
                    "price": price,
                    "stock": product[5] if len(product) > 5 else None,
                    "price_includes_igv": product[6] if len(product) > 6 else 0
                }
                
                # Actualizar UI de selección
                stock_msg = SalesUIHelpers.check_and_display_stock(self.app, product)
                
                # Actualizar campos editables
                if hasattr(self.app, 'entry_custom_price'):
                    self.app.entry_custom_price.delete(0, tk.END)
                    self.app.entry_custom_price.insert(0, f"{price:.2f}")
                
                # Poner foco en cantidad
                self.app.entry_quantity.delete(0, tk.END)
                self.app.entry_quantity.insert(0, "1")
                self.app.entry_quantity.focus()
                self.app.entry_quantity.select_range(0, tk.END)
                
        except Exception as e:
            print(f"Error al seleccionar producto: {e}")
            self.app.selected_product = None

    def show_product_selection_window(self, products):
        # Nota: Este método parece ser usado cuando hay múltiples coincidencias en una búsqueda manual (no implementada completamente en el código original mostrado, pero estaba ahí)
        # El código original tenía un error: usaba 'tree' que no estaba definido. Lo corregiré aquí.
        
        win = tb.Toplevel(master=self.app)
        self.app.begin_modal_construction(win)
        win.title("Seleccionar producto")
        win.geometry("600x400")
        self.app.set_app_icon(win)

        cols = ("id", "code", "name", "unit", "price")
        tree = tb.Treeview(win, columns=cols, show="headings", bootstyle="info")
        tree.heading("id", text="ID")
        tree.heading("code", text="Código")
        tree.heading("name", text="Nombre")
        tree.heading("unit", text="Unidad")
        tree.heading("price", text="Precio")
        tree.pack(fill="both", expand=True, padx=10, pady=10)

        SalesUIHelpers.bind_mousewheel(tree)
        try:
            SalesUIHelpers.bind_mousewheel_to_children(win)
        except Exception:
            pass

        for p in products:
            tree.insert("", "end", values=p)

        def select():
            sel = tree.focus()
            if not sel:
                self.app.show_warning("Atención", "Seleccione un producto.")
                return
            values = tree.item(sel, "values")
            product_id, code, name, unit, price = values
            # Llamar a append_item de la app principal
            self.app.sales_ui.append_item(product_id, name, unit, float(price))
            win.destroy()

        btn = tb.Button(win, text="Agregar al comprobante", bootstyle="success-outline", command=select)
        btn.pack(pady=5)

        win.update_idletasks()
        self.app.center_window(win, 600, 400)
        self.app.reveal_modal(win)
        self.app.make_modal(win)

    def show_product_form_window(self, prefill_name=""):
        win = tb.Toplevel(master=self.app)
        self.app.begin_modal_construction(win)
        win.title("Registrar producto")
        win.geometry("400x300")
        self.app.set_app_icon(win)

        lf = tb.Labelframe(win, text="Nuevo Producto", padding=10, bootstyle="info")
        lf.pack(fill="both", expand=True, padx=10, pady=10)
        lf.columnconfigure(1, weight=1)

        try:
            SalesUIHelpers.bind_mousewheel_to_children(lf)
        except Exception:
            pass

        tb.Label(lf, text="Nombre:").grid(row=0, column=0, sticky="e", padx=5, pady=5)
        entry_name = tb.Entry(lf)
        entry_name.grid(row=0, column=1, sticky="we", padx=5, pady=5)

        tb.Label(lf, text="Unidad (ej. unidad, kg):").grid(row=1, column=0, sticky="e", padx=5, pady=5)
        entry_unit = tb.Entry(lf)
        entry_unit.grid(row=1, column=1, sticky="we", padx=5, pady=5)

        tb.Label(lf, text="Precio unitario:").grid(row=2, column=0, sticky="e", padx=5, pady=5)
        entry_price = tb.Entry(lf)
        entry_price.grid(row=2, column=1, sticky="we", padx=5, pady=5)

        entry_name.insert(0, prefill_name)
        
        # Forzar mayúsculas
        self.app.setup_uppercase_entry(entry_name, entry_unit)

        def save_product():
            try:
                product_id = self.app.product_service.create_product(
                    entry_name.get().strip(),
                    entry_unit.get().strip(),
                    entry_price.get().strip(),
                    code=None
                )
                product = self.app.product_service.get_product_by_id(product_id)
                # product: (id, code, name, unit, price, stock, ...)
                _, code, name, unit, price = product[:5]
                self.app.sales_ui.append_item(product_id, name, unit, float(price))
                win.destroy()
            except ValueError as e:
                self.app.show_error("Error", str(e))

        btn_save = tb.Button(lf, text="Guardar producto", bootstyle="success-outline", command=save_product)
        btn_save.grid(row=3, column=0, columnspan=2, pady=10)

        win.update_idletasks()
        self.app.center_window(win, 400, 300)
        self.app.reveal_modal(win)
        self.app.make_modal(win)

    # -------------------------------------------------------------------------
    # GESTIÓN DE PRODUCTOS (PESTAÑA ACTUALIZAR INFO)
    # -------------------------------------------------------------------------

    def filter_products(self, event=None):
        """Filtra productos en tiempo real."""
        # Verificar que el caché exista
        if not hasattr(self.app, 'all_products_cache'):
            self.app.all_products_cache = []
            return
            
        query = self.app.entry_search_product.get().lower()
        if query == "buscar producto...":
            query = ""

        tree = self.app.tree_products
        tree.delete(*tree.get_children())

        # Configurar colores de stock si aún no se hizo
        if not getattr(self, "_stock_tags_configured", False):
            tree.tag_configure("stock_empty", background="#FDE4E4", foreground="#5A0B0B")
            tree.tag_configure("stock_low", background="#FFF3CD", foreground="#665200")
            tree.tag_configure("stock_ok", background="", foreground="")
            tree.tag_configure("stock_na", background="", foreground="")
            self._stock_tags_configured = True

        for p in self.app.all_products_cache:
            code = str(p[1]).lower() if p[1] else ""
            name = str(p[2]).lower()

            if query in code or query in name:
                stock = p[5] if len(p) > 5 else None
                tag = self._resolve_stock_tag(stock)
                stock_display = self._format_stock_display(stock)
                price_value = p[4]
                try:
                    price_display = f"{float(price_value):.2f}"
                except (TypeError, ValueError):
                    price_display = str(price_value or "0.00")

                tree.insert(
                    "",
                    "end",
                    values=(p[0], p[1], p[2], p[3], price_display, stock_display),
                    tags=(tag,)
                )

    def on_product_select(self, event):
        """Carga datos del producto seleccionado en el formulario."""
        sel = self.app.tree_products.focus()
        if not sel:
            return
        values = self.app.tree_products.item(sel, "values")
        # values: (id, code, name, unit, price)
        
        self.app.var_prod_id.set(values[0])
        self.app.var_prod_code.set(values[1] if values[1] != 'None' else '')
        self.app.var_prod_name.set(values[2])
        self.app.var_prod_unit.set(values[3])
        self.app.var_prod_price.set(values[4])

        # Si el usuario selecciona algo, estamos en modo edición
        setattr(self.app, "_force_new_product", False)
        
        # Buscar stock real en cache
        product = next((p for p in self.app.all_products_cache if str(p[0]) == str(values[0])), None)
        if product and len(product) > 5:
            stock = product[5]
            if stock is None:
                self.app.var_prod_stock.set("")
            else:
                try:
                    stock_val = float(stock)
                    if stock_val.is_integer():
                        self.app.var_prod_stock.set(str(int(stock_val)))
                    else:
                        self.app.var_prod_stock.set(f"{stock_val:.2f}")
                except (TypeError, ValueError):
                    self.app.var_prod_stock.set(str(stock))
        else:
            self.app.var_prod_stock.set("")

    def clear_product_form(self):
        """Limpia el formulario de productos."""
        self.app.var_prod_id.set("")
        self.app.var_prod_code.set("")
        self.app.var_prod_name.set("")
        self.app.var_prod_unit.set("")
        self.app.var_prod_price.set("")
        self.app.var_prod_stock.set("")
        try:
            self.app.tree_products.selection_remove(self.app.tree_products.selection())
            self.app.tree_products.focus("")
        except Exception:
            pass

        # Próximo guardado fuerza creación
        setattr(self.app, "_force_new_product", True)

    def save_product_action(self):
        """Guarda (crea o actualiza) un producto desde la pestaña de gestión."""
        p_id = self.app.var_prod_id.get()
        code = self.app.var_prod_code.get().strip() or None
        name = self.app.var_prod_name.get().strip()
        unit = self.app.var_prod_unit.get().strip()
        price_str = self.app.var_prod_price.get().strip()
        stock_str = self.app.var_prod_stock.get().strip()

        if not name or not price_str:
            self.app.show_warning("Atención", "Nombre y Precio son obligatorios.")
            return

        try:
            force_new = bool(getattr(self.app, "_force_new_product", False))

            if p_id and not force_new:
                # Actualizar
                self.app.product_service.update_product(int(p_id), name, unit, price_str, code, stock=stock_str)
                self.app.show_success("Éxito", "Producto actualizado.")
            else:
                # Crear
                self.app.product_service.create_product(name, unit, price_str, code, stock=stock_str)
                self.app.show_success("Éxito", "Producto creado.")
            
            self.clear_product_form()
            setattr(self.app, "_force_new_product", False)
            # Recargar caché y tabla
            self.app.all_products_cache = self.app.product_service.search_products("")
            self.filter_products()
            
        except ValueError as e:
            self.app.show_error("Error", str(e))

    # --- Helpers de stock ---
    def _resolve_stock_tag(self, stock):
        if stock is None:
            return "stock_na"
        try:
            stock_val = float(stock)
        except (TypeError, ValueError):
            return "stock_na"
        if stock_val <= 0:
            return "stock_empty"
        if stock_val <= 5:
            return "stock_low"
        return "stock_ok"

    def _format_stock_display(self, stock):
        if stock is None:
            return "Sin control"
        try:
            if float(stock).is_integer():
                return str(int(float(stock)))
            return f"{float(stock):.2f}"
        except (TypeError, ValueError):
            return str(stock)

    def delete_product_action(self):
        """Elimina el producto seleccionado."""
        p_id = self.app.var_prod_id.get()
        if not p_id:
            self.app.show_warning("Atención", "Seleccione un producto para eliminar.")
            return

        if not self.app.show_question("Confirmar", "¿Está seguro de eliminar este producto?"):
            return

        try:
            # Guardar snapshot para deshacer
            sel = self.app.tree_products.focus() or (self.app.tree_products.selection()[0] if self.app.tree_products.selection() else "")
            values = self.app.tree_products.item(sel, "values") if sel else None
            if values:
                # (id, code, name, unit, price, stock_display)
                self.app.undo_last_product = tuple(values)

            self.app.product_service.delete_product(int(p_id))
            self.app.show_success("Éxito", "Producto eliminado.")
            self.clear_product_form()
            self.app.all_products_cache = self.app.product_service.search_products("")
            self.filter_products()
        except Exception as e:
            self.app.show_error("Error", f"No se pudo eliminar: {e}")

    def import_products_from_excel(self):
        """Importar productos desde un archivo Excel."""
        if load_workbook is None:
            self.app.show_error("Error", "La librería 'openpyxl' no está instalada.")
            return

        file_path = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=[("Excel Files", "*.xlsx;*.xls")]
        )
        
        if not file_path:
            return

        try:
            wb = load_workbook(file_path)
            ws = wb.active

            header_rows = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            if not header_rows:
                self.app.show_warning("Atención", "El archivo está vacío.")
                return

            headers = [str(h).strip().lower() if h is not None else "" for h in header_rows[0]]

            def _find_col(*aliases):
                for name in aliases:
                    if name in headers:
                        return headers.index(name)
                return None

            col_code = _find_col("codigo", "código", "code")
            col_name = _find_col("nombre", "name")
            col_unit = _find_col("unidad", "unit")
            col_price = _find_col("precio", "price", "precio unitario")
            col_stock = _find_col("stock", "existencia", "cantidad", "inventario")

            # Compatibilidad con archivos antiguos sin cabeceras esperadas
            if col_code is None:
                col_code = 0
            if col_name is None:
                col_name = 1
            if col_unit is None:
                col_unit = 2
            if col_price is None:
                col_price = 3

            def _val(row, idx):
                return row[idx] if idx is not None and idx < len(row) else None

            count = 0
            errors = 0

            for row in ws.iter_rows(min_row=2, values_only=True):
                name = _val(row, col_name)
                if not name:
                    continue

                price = _val(row, col_price)
                if price is None or str(price).strip() == "":
                    errors += 1
                    continue

                code = _val(row, col_code)
                unit = _val(row, col_unit)
                stock = _val(row, col_stock)

                name_norm = str(name).strip().upper()
                unit_norm = str(unit).strip().upper() if unit else "UNI"
                code_norm = str(code).strip() if code else None

                try:
                    self.app.product_service.create_product(
                        name_norm,
                        unit_norm,
                        str(price).strip(),
                        code_norm,
                        stock=stock,
                    )
                    count += 1
                except Exception:
                    errors += 1

            msg = f"Se importaron {count} productos correctamente."
            if errors > 0:
                msg += f"\nHubo {errors} errores."

            self.app.show_success("Importación", msg)

            # Recargar
            self.app.all_products_cache = self.app.product_service.search_products("")
            self.filter_products()

        except Exception as e:
            self.app.show_error("Error de Importación", str(e))
